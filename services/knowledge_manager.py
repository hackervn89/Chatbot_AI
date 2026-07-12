"""
Knowledge Manager — Quản lý tài liệu tri thức (CRUD, Versioning, Re-indexing).
"""
import os
import re
import time
import zipfile
from xml.etree import ElementTree
from sqlalchemy.orm import Session
from sqlalchemy.sql import func
from pypdf import PdfReader

from config import (
    IMAGES_DIR, CATEGORY_CORE, CATEGORY_UPDATABLE, CATEGORY_PERSONAL
)
from database import get_db_session
from models import Document, DocumentVersion, KnowledgeChunk, ImageMapping
from services.rag_pipeline import semantic_chunk, get_embedding
from services.audit_logger import log_document_action


# ==================== FILE PARSERS ====================

def parse_pdf(file_path: str) -> str:
    """Trích xuất văn bản từ PDF"""
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def parse_docx(file_path: str) -> tuple:
    """Trích xuất văn bản + hình ảnh từ DOCX"""
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    
    full_text = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    raw_text = "\n".join(full_text)
    
    # Trích xuất hình ảnh
    image_mappings = _extract_docx_images(file_path)
    
    return raw_text, image_mappings


def _extract_docx_images(docx_path: str) -> dict:
    """Bóc tách hình ảnh từ file DOCX (ZIP format)"""
    safe_name = re.sub(r'[^a-zA-Z0-9]', '_', os.path.basename(docx_path).replace('.docx', ''))
    doc_image_dir = os.path.join(IMAGES_DIR, safe_name)
    os.makedirs(doc_image_dir, exist_ok=True)
    
    mappings = {}
    try:
        with zipfile.ZipFile(docx_path) as z:
            rels_xml = z.read('word/_rels/document.xml.rels')
            rels_root = ElementTree.fromstring(rels_xml)
            
            id_to_file = {}
            for child in rels_root:
                r_id = child.attrib.get('Id')
                target = child.attrib.get('Target')
                if r_id and target and 'media' in target:
                    id_to_file[r_id] = os.path.basename(target)
            
            doc_xml = z.read('word/document.xml')
            doc_root = ElementTree.fromstring(doc_xml)
            
            counter = 1
            for elem in doc_root.iter():
                if elem.tag.endswith('drawing'):
                    blips = [e for e in elem.iter() if e.tag.endswith('blip')]
                    if blips:
                        ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
                        r_id = blips[0].attrib.get(f'{ns}embed')
                        if r_id and r_id in id_to_file:
                            filename = id_to_file[r_id]
                            ext = os.path.splitext(filename)[1] or '.png'
                            dest = f"hinh_{counter}{ext}"
                            dest_path = os.path.join(doc_image_dir, dest)
                            try:
                                with open(dest_path, 'wb') as f:
                                    f.write(z.read(f"word/media/{filename}"))
                                mappings[f"Hình {counter}"] = f"images/{safe_name}/{dest}"
                                counter += 1
                            except Exception as ex:
                                print(f"[KM] Lỗi trích xuất ảnh: {ex}")
    except Exception as e:
        print(f"[KM] Warning: Không thể trích xuất ảnh DOCX: {e}")
    
    return mappings


def parse_srt(file_path: str) -> str:
    """Làm sạch file phụ đề SRT"""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    content = re.sub(r'\d{2}:\d{2}:\d{2}[,.]\d{3}\s*-->\s*\d{2}:\d{2}:\d{2}[,.]\d{3}', '', content)
    lines = [l.strip() for l in content.split('\n') if l.strip() and not l.strip().isdigit()]
    return "\n".join(lines)


def parse_text(file_path: str) -> str:
    """Đọc file text/markdown"""
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()


# ==================== KNOWLEDGE CRUD ====================

def ingest_document(
    file_path: str,
    title: str = None,
    category: str = CATEGORY_CORE,
    description: str = "",
    created_by: str = "admin",
    status: str = "active",
    db: Session = None
) -> dict:
    """
    Nạp tài liệu mới vào hệ thống RAG (hỗ trợ lưu nháp hoặc hoạt động ngay).
    
    Returns: {"status": "success/error", "message": "...", "details": {...}}
    """
    if not os.path.exists(file_path):
        return {"status": "error", "message": "File không tồn tại."}

    should_close = False
    if db is None:
        db = get_db_session()
        should_close = True

    filename = os.path.basename(file_path)
    file_ext = os.path.splitext(filename)[1].lower()
    
    if not title:
        title = os.path.splitext(filename)[0].replace('_', ' ')

    # Source name chuẩn hóa
    doc_source = f"{category}/{filename}"

    try:
        # 1. Parse file
        print(f"[KM] Đang xử lý: {filename} ({file_ext})")
        raw_text = ""
        image_mappings = {}

        if file_ext == '.pdf':
            raw_text = parse_pdf(file_path)
        elif file_ext == '.docx':
            raw_text, image_mappings = parse_docx(file_path)
        elif file_ext == '.srt':
            raw_text = parse_srt(file_path)
        elif file_ext in ('.txt', '.md'):
            raw_text = parse_text(file_path)
        else:
            return {"status": "error", "message": f"Định dạng {file_ext} không được hỗ trợ."}

        if not raw_text.strip():
            return {"status": "error", "message": "Tài liệu trống."}

        # 2. Kiểm tra tài liệu đã tồn tại chưa
        doc = db.query(Document).filter(Document.source == doc_source).first()
        
        if doc:
            # Tạo version mới
            doc.current_version += 1
            version = DocumentVersion(
                document_id=doc.id,
                version=doc.current_version,
                raw_text=raw_text,
                change_summary=f"Cập nhật từ file {filename}",
                changed_by=created_by
            )
            db.add(version)
            
            # Xóa chunks cũ
            db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == doc.id).delete()
            db.query(ImageMapping).filter(ImageMapping.document_id == doc.id).delete()
            
            doc.title = title
            doc.raw_text = raw_text
            doc.description = description
            doc.updated_at = func.now()
            action = "UPDATE"
        else:
            # Tạo tài liệu mới
            doc = Document(
                title=title,
                source=doc_source,
                category=category,
                description=description,
                file_type=file_ext.lstrip('.'),
                raw_text=raw_text,
                created_by=created_by
            )
            db.add(doc)
            db.commit()
            db.refresh(doc)
            
            # Tạo version đầu tiên
            version = DocumentVersion(
                document_id=doc.id,
                version=1,
                raw_text=raw_text,
                change_summary="Phiên bản gốc",
                changed_by=created_by
            )
            db.add(version)
            action = "CREATE"

        db.commit()
        db.refresh(doc)

        # Nếu là tài liệu nháp (draft), chỉ lưu file/image và không tạo chunks/embeddings ngay
        if status == "draft":
            doc.status = "draft"
            doc.chunk_count = 0
            doc.is_active = False  # Chưa kích hoạt cho RAG
            db.commit()
            
            # Lưu image mappings
            img_count = 0
            for hinh_key, img_path in image_mappings.items():
                db.add(ImageMapping(document_id=doc.id, hinh_key=hinh_key, img_rel_path=img_path))
                img_count += 1
            db.commit()
            
            # Audit log
            log_document_action("CREATE_DRAFT" if action == "CREATE" else "UPDATE_DRAFT", doc.id, actor=created_by, details={
                "title": title, "category": category, "file_type": file_ext
            }, db=db)
            
            print(f"[KM] Đã tạo nháp tài liệu {title} (ID={doc.id}). Đợi duyệt.")
            return {
                "status": "success",
                "message": "Nạp tài liệu nháp thành công! Vui lòng duyệt tại Dashboard.",
                "details": {
                    "doc_id": doc.id,
                    "title": title,
                    "chunks": 0,
                    "images": img_count,
                    "version": doc.current_version,
                    "status": "draft"
                }
            }

        # Nếu nạp trực tiếp (active)
        doc.status = "active"
        db.commit()

        # 3. Semantic chunking
        chunks = semantic_chunk(raw_text)
        print(f"[KM] Đã chia thành {len(chunks)} chunks.")

        # 4. Generate embeddings & save chunks bằng BATCH để tránh rate-limit
        from services.rag_pipeline import get_embeddings_batch
        
        batch_size = 50
        inserted = 0
        
        for idx in range(0, len(chunks), batch_size):
            chunk_batch = chunks[idx : idx + batch_size]
            batch_texts = [c["text"] for c in chunk_batch]
            
            print(f"[KM] Đang sinh batch embedding cho {len(batch_texts)} chunks (từ {idx} đến {idx + len(batch_texts)})...")
            vectors = get_embeddings_batch(batch_texts)
            
            for j, chunk_data in enumerate(chunk_batch):
                vector = vectors[j] if j < len(vectors) else [0.0] * 3072
                chunk_obj = KnowledgeChunk(
                    document_id=doc.id,
                    chunk_index=idx + j,
                    text=chunk_data["text"],
                    embedding=vector,
                    chunk_metadata=chunk_data.get("metadata", {})
                )
                db.add(chunk_obj)
                inserted += 1
                
            db.commit()
            time.sleep(1.0)  # Sleep 1s giữa các batch để an toàn cho Free Tier

        # 5. Save image mappings
        img_count = 0
        for hinh_key, img_path in image_mappings.items():
            db.add(ImageMapping(document_id=doc.id, hinh_key=hinh_key, img_rel_path=img_path))
            img_count += 1

        # Update chunk count
        doc.chunk_count = inserted
        db.commit()

        # 6. Audit log
        log_document_action(action, doc.id, actor=created_by, details={
            "title": title, "chunks": inserted, "images": img_count,
            "category": category, "file_type": file_ext
        }, db=db)

        print(f"[KM] Nạp thành công! {inserted} chunks, {img_count} ảnh.")
        return {
            "status": "success",
            "message": f"Nạp tài liệu thành công!",
            "details": {
                "doc_id": doc.id,
                "title": title,
                "chunks": inserted,
                "images": img_count,
                "version": doc.current_version
            }
        }

    except Exception as e:
        db.rollback()
        print(f"[KM] Lỗi nạp tài liệu: {e}")
        return {"status": "error", "message": str(e)}
    finally:
        if should_close:
            db.close()


def toggle_document(doc_id: int, is_active: bool, actor: str = "admin", db: Session = None) -> bool:
    """Bật/tắt tài liệu khỏi RAG (không xóa)"""
    should_close = db is None
    if db is None:
        db = get_db_session()
    try:
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if doc:
            doc.is_active = is_active
            db.commit()
            status = "ACTIVATE" if is_active else "DEACTIVATE"
            log_document_action(status, doc_id, actor=actor, db=db)
            return True
        return False
    finally:
        if should_close:
            db.close()


def delete_document(doc_id: int, actor: str = "admin", db: Session = None) -> bool:
    """Xóa tài liệu và tất cả chunks liên quan"""
    should_close = db is None
    if db is None:
        db = get_db_session()
    try:
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if doc:
            title = doc.title
            db.delete(doc)  # CASCADE sẽ xóa chunks, versions, images
            db.commit()
            log_document_action("DELETE", doc_id, actor=actor, details={"title": title}, db=db)
            return True
        return False
    finally:
        if should_close:
            db.close()


def get_all_documents(db: Session, category: str = None, active_only: bool = False) -> list:
    """Lấy danh sách tài liệu"""
    query = db.query(Document)
    if category:
        query = query.filter(Document.category == category)
    if active_only:
        query = query.filter(Document.is_active == True)
    return query.order_by(Document.updated_at.desc()).all()


def get_document_detail(db: Session, doc_id: int) -> dict:
    """Lấy chi tiết tài liệu kèm versions và chunks"""
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        return None
    
    versions = db.query(DocumentVersion).filter(
        DocumentVersion.document_id == doc_id
    ).order_by(DocumentVersion.version.desc()).all()
    
    chunks = db.query(KnowledgeChunk).filter(
        KnowledgeChunk.document_id == doc_id
    ).order_by(KnowledgeChunk.chunk_index).all()
    
    return {
        "document": doc,
        "versions": versions,
        "chunks": chunks,
        "chunk_count": len(chunks)
    }


def publish_document(doc_id: int, actor: str = "admin", db: Session = None) -> dict:
    """Kích hoạt tài liệu nháp: Cắt chunks → Tạo vector embeddings → Lưu CSDL"""
    should_close = db is None
    if db is None:
        db = get_db_session()
        should_close = True
        
    try:
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if not doc:
            return {"status": "error", "message": "Không tìm thấy tài liệu."}
            
        if doc.status == "active" and doc.chunk_count > 0:
            return {"status": "error", "message": "Tài liệu này đã hoạt động rồi."}
            
        print(f"[KM] Đang phát hành tài liệu: {doc.title} (ID={doc.id})")
        
        # 1. Xóa các chunks cũ nếu có
        db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == doc.id).delete()
        
        # 2. Cắt chunks ngữ nghĩa
        chunks = semantic_chunk(doc.raw_text)
        
        # 3. Tạo vector embeddings bằng BATCH để tránh rate-limit
        from services.rag_pipeline import get_embeddings_batch
        from config import EMBEDDING_DIMENSION
        
        batch_size = 50
        inserted = 0
        
        for idx in range(0, len(chunks), batch_size):
            chunk_batch = chunks[idx : idx + batch_size]
            batch_texts = [c["text"] for c in chunk_batch]
            vectors = get_embeddings_batch(batch_texts)
            
            for j, chunk_data in enumerate(chunk_batch):
                vector = vectors[j] if j < len(vectors) else [0.0] * EMBEDDING_DIMENSION
                chunk_obj = KnowledgeChunk(
                    document_id=doc.id,
                    chunk_index=idx + j,
                    text=chunk_data["text"],
                    embedding=vector,
                    chunk_metadata=chunk_data.get("metadata", {})
                )
                db.add(chunk_obj)
                inserted += 1
                
            db.commit()
            time.sleep(1.0)  # Sleep 1s giữa các batch
            
        # 4. Cập nhật trạng thái
        doc.status = "active"
        doc.is_active = True
        doc.chunk_count = inserted
        doc.updated_at = func.now()
        db.commit()
        
        # 5. Ghi Audit log
        log_document_action("PUBLISH", doc.id, actor=actor, details={
            "title": doc.title, "chunks": inserted
        }, db=db)
        
        return {
            "status": "success",
            "message": f"Kích hoạt tài liệu thành công! Đã nạp {inserted} chunks vào hệ thống RAG.",
            "details": {"doc_id": doc.id, "chunks": inserted}
        }
        
    except Exception as e:
        db.rollback()
        print(f"[KM] Lỗi kích hoạt tài liệu: {e}")
        return {"status": "error", "message": str(e)}
    finally:
        if should_close:
            db.close()


def analyze_document_draft(doc_id: int, db: Session = None) -> dict:
    """Gọi AI phân tích tài liệu nháp: tóm tắt, FAQ, phát hiện mâu thuẫn tri thức"""
    should_close = db is None
    if db is None:
        db = get_db_session()
        should_close = True
        
    try:
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if not doc:
            return {"status": "error", "message": "Không tìm thấy tài liệu."}
            
        # Lấy một số tài liệu hoạt động khác cùng mảng để so sánh
        other_docs = db.query(Document).filter(
            Document.id != doc.id,
            Document.is_active == True,
            Document.category == doc.category
        ).limit(3).all()
        
        other_contexts = ""
        if other_docs:
            other_contexts = "\n\nDưới đây là một số tài liệu tri thức cũ đang hoạt động trong hệ thống để đối chiếu:\n"
            for od in other_docs:
                other_contexts += f"- Tiêu đề: {od.title}\n  Nội dung tóm lược: {od.raw_text[:1000]}...\n\n"
                
        # Build prompt
        system_prompt = (
            "Bạn là chuyên viên phân tích chất lượng dữ liệu tri thức của hệ thống RAG thuộc Khối Đảng. "
            "Nhiệm vụ của bạn là đọc tài liệu mới gửi đến và đối chiếu với các tài liệu cũ "
            "nhằm phát hiện các mâu thuẫn nghiệp vụ hoặc chồng chéo (giờ giấc, quy trình, quyền hạn)."
        )
        
        user_message = (
            f"Tài liệu mới cần phân tích:\n"
            f"Tiêu đề: {doc.title}\n"
            f"Nội dung:\n{doc.raw_text[:4000]}\n"
            f"{other_contexts}\n"
            f"Hãy đưa ra báo cáo phân tích theo cấu trúc sau (viết bằng định dạng Markdown tiếng Việt):\n"
            f"1. **Tóm tắt nội dung chính**: (2-3 câu ngắn gọn)\n"
            f"2. **Phát hiện mâu thuẫn tri thức**: (Chỉ rõ điểm mâu thuẫn/chồng chéo hoặc ghi 'Không phát hiện mâu thuẫn' nếu tri thức đồng bộ)\n"
            f"3. **Gợi ý 3 câu hỏi FAQ mẫu kèm câu trả lời ngắn**."
        )
        
        from services.ai_engine import call_ai
        reply, model_name, _ = call_ai(
            system_prompt=system_prompt,
            user_message=user_message,
            temperature=0.3
        )
        
        if not reply:
            reply = "Không thể gọi AI phân tích tại thời điểm này. Vui lòng kiểm tra lại cấu hình API Key."
            
        return {
            "status": "success",
            "analysis": reply,
            "model": model_name
        }
        
    except Exception as e:
        print(f"[KM] Lỗi AI phân tích: {e}")
        return {"status": "error", "message": str(e)}
    finally:
        if should_close:
            db.close()
